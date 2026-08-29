"""Task 8: add two liquefaction drive-type parameters to the Data Inputs workbook.

Both come from the British Columbia Environmental Assessment Office's assessment
of Ksi Lisims LNG. They are **sensitivity** parameters: the central liquefaction
factor stays 0.29 for every terminal and the 0.29 assertion in build_results.py
is untouched.

The boundary note matters and is written into both rows: the two figures are
**facility total intensity including marine sources**, not liquefaction alone,
so they are not like-for-like with the 0.29 liquefaction-stage factor. The
sensitivity says so wherever it reports them.

Run once, from the repository root:

    python tools/add_liquefaction_drive_params.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

BOOK = Path(__file__).resolve().parents[1] / "Inputs" / "Canada_LNG_Data_Inputs.xlsx"

SOURCE = (
    "British Columbia Environmental Assessment Office, Assessment Report for "
    "Ksi Lisims LNG, 7 August 2025, Canadian Impact Assessment Registry "
    "document 163192E, pages 847 to 848"
)
URL = "https://iaac-aeic.gc.ca/050/documents/p82797/163192E.pdf"
BOUNDARY_NOTE = (
    "SENSITIVITY ONLY - the central liquefaction factor stays 0.29 for every "
    "terminal. BOUNDARY: this is the facility total intensity including marine "
    "sources, not the liquefaction stage alone, so it is not like for like with "
    "the 0.29 liquefaction factor on the Emission Factors sheet. Any comparison "
    "against 0.29 is therefore approximate and is labelled as such wherever it "
    "is reported."
)

ROWS = [
    {
        "parameter": "liquefaction_electric_drive_gas_generation",
        "value": 0.156,
        "unit": "tCO2e per t LNG",
        "type": "cited",
        "source": f"{SOURCE} - Alternative Case, gas-fired power barges",
        "source_url": URL,
        "vintage": "2025",
        "notes": BOUNDARY_NOTE,
    },
    {
        "parameter": "liquefaction_electric_drive_grid",
        "value": 0.021,
        "unit": "tCO2e per t LNG",
        "type": "cited",
        "source": f"{SOURCE} - Base Case, grid supply",
        "source_url": URL,
        "vintage": "2025",
        "notes": BOUNDARY_NOTE,
    },
]


def main() -> None:
    wb = openpyxl.load_workbook(BOOK)
    ws = wb["Parameters"]
    headers = {
        str(ws.cell(row=1, column=c).value).strip(): c
        for c in range(1, ws.max_column + 1)
    }
    for need in ("parameter", "value", "unit", "type", "source", "source_url"):
        if need not in headers:
            raise KeyError(f"Parameters sheet has no '{need}' column.")
    existing = {
        str(ws.cell(row=r, column=headers["parameter"]).value).strip()
        for r in range(2, ws.max_row + 1)
    }
    row = ws.max_row + 1
    added = []
    for spec in ROWS:
        if spec["parameter"] in existing:
            print(f"skip (already present): {spec['parameter']}")
            continue
        for key, value in spec.items():
            if key in headers:
                ws.cell(row=row, column=headers[key], value=value)
        added.append(spec["parameter"])
        row += 1
    if added:
        wb.save(BOOK)
        print(f"added to Parameters: {', '.join(added)}")
    else:
        print("nothing to add")


if __name__ == "__main__":
    main()
