"""lifespan_sensitivity_low 30 years: record the document, page and quotation.

One-off. No value changes: 30 stays 30. The model never writes Inputs/; this
script does, once.

The row already said the 30 years "matches the operating life Summit Lake PG LNG
states in its impact assessment", with `source_url` reading "Not available". The
document was found on 3 September 2026 and opened; the sentence is on page 11 of
106, under the heading "Estimates and Purpose":

    "The Project is expected to operate for 30 years."

The same document states it a second time on page 52, section 5 Decommissioning:
"The decommissioning process for the Project is set to commence after a minimum
of 30 years after the anticipated phase 1 construction year."

Type moves `assumption` -> `cited`, matching how `lifespan_sensitivity_high` = 50
is already typed: the bound itself now rests on a named document, even though the
choice to use it as the sensitivity low remains a modelling judgement.

Two honest notes on the citation:

1. The BC EAO registry stores the file under the name "Draft Initial Project
   Description.pdf", but the document's own title page reads "Initial Project
   Description", dated 12 February 2024. The registry filename is recorded here
   so nobody thinks the link is to a different document.
2. The Impact Assessment Agency of Canada suspended the impact assessment
   timeline on 19 July 2024 at the proponent's request. The stated operating life
   is unaffected - it is the proponent's own design life - but a reviewer should
   know the assessment is paused rather than complete.

Run from the repository root:

    python tools/cite_lifespan_sensitivity_low.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "Inputs" / "Canada_LNG_Data_Inputs.xlsx"

PARAM = "lifespan_sensitivity_low"

EAO_PDF = (
    "https://www.projects.eao.gov.bc.ca/api/public/document/"
    "65cce8f768c392001b2d5692/download/Draft%20Initial%20Project%20Description.pdf"
)
EAO_PROJECT = "https://projects.eao.gov.bc.ca/p/65c661a8399db00022d48849/documents"
IAAC_SUMMARY_PDF = "https://ceaa.gc.ca/050/documents/p87307/155565E.pdf"
IAAC_PROJECT = "https://iaac-aeic.gc.ca/050/evaluations/proj/87307?culture=en-CA"

SOURCE = (
    "Lower bound for the lifespan sensitivity, and now a cited one. JX LNG "
    "Canada Ltd., Summit Lake PG LNG Project - Liquified Natural Gas Facility, "
    "INITIAL PROJECT DESCRIPTION (BCEAA 2018, IAA 2019), prepared by Keywest "
    "Projects Ltd., 12 February 2024. PAGE 11 OF 106, under the heading "
    "'Estimates and Purpose': 'The Project is expected to operate for 30 "
    "years.' Stated again on page 52, section 5 Decommissioning: 'The "
    "decommissioning process for the Project is set to commence after a "
    "minimum of 30 years after the anticipated phase 1 construction year.' "
    "The same sentence appears on page 6 of 17 of the English Summary filed "
    "with the Impact Assessment Agency of Canada, " + IAAC_SUMMARY_PDF + " . "
    "NOTE ON THE FILENAME: the British Columbia Environmental Assessment "
    "Office registry stores the file as 'Draft Initial Project "
    "Description.pdf' while the document's own title page reads 'Initial "
    "Project Description'; the link is to that document, not a different one. "
    "Registry landing page: " + EAO_PROJECT + " ; federal registry: "
    + IAAC_PROJECT + " . NOTE ON STATUS: the Agency suspended the impact "
    "assessment timeline on 19 July 2024 at the proponent's request, so the "
    "assessment is paused rather than complete. The 30 years is the "
    "proponent's own stated design life and does not depend on the assessment "
    "concluding. Located and opened 3 September 2026; source_url was 'Not "
    "available' until then."
)

NOTES = (
    "About 25 per cent below the 40 year case. The 30 years is one proponent's "
    "stated design life, not an industry norm: the four main export projects "
    "in this register run on 40 year CER export licence terms "
    "(lifecycle_years_default), and the CER may issue licences of up to 50 "
    "years (lifespan_sensitivity_high). The three together bracket the life "
    "assumption from a real short case to a real regulatory ceiling."
)


def _headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(row=1, column=c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=c).value is not None
    }


def main() -> None:
    wb = openpyxl.load_workbook(DATA)
    ws = wb["Parameters"]
    headers = _headers(ws)

    row = None
    for r in range(2, ws.max_row + 1):
        value = ws.cell(row=r, column=headers["parameter"]).value
        if value is not None and str(value).strip() == PARAM:
            row = r
            break
    if row is None:
        raise SystemExit(f"{PARAM} not found on the Parameters sheet")

    before = {
        h: ws.cell(row=row, column=c).value
        for h, c in headers.items()
        if h in ("value", "type", "source_url", "vintage")
    }

    assert float(ws.cell(row=row, column=headers["value"]).value) == 30.0, (
        "value is not 30; refusing to write a citation onto a different number"
    )

    ws.cell(row=row, column=headers["type"], value="cited")
    ws.cell(row=row, column=headers["source"], value=SOURCE)
    ws.cell(row=row, column=headers["source_url"], value=EAO_PDF)
    ws.cell(row=row, column=headers["vintage"], value=2024)
    ws.cell(row=row, column=headers["notes"], value=NOTES)

    after = {
        h: ws.cell(row=row, column=c).value
        for h, c in headers.items()
        if h in ("value", "type", "source_url", "vintage")
    }

    wb.save(DATA)
    for key in ("value", "type", "source_url", "vintage"):
        print(f"{key}: {before[key]!r} -> {after[key]!r}")
    print("saved")


if __name__ == "__main__":
    main()
