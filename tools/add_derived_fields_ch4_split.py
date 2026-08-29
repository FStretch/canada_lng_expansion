"""Task 1: record the CO2/CH4 split construction on the Data Inputs workbook.

Adds three rows to the `Derived Fields` sheet. No emission factor, parameter or
register value changes. Run once, from the repository root:

    python tools/add_derived_fields_ch4_split.py

The model never writes to Inputs/. This script is the deliberate, recorded edit.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

BOOK = Path(__file__).resolve().parents[1] / "Inputs" / "Canada_LNG_Data_Inputs.xlsx"
SHEET = "Derived Fields"

ROWS = [
    (
        "upstream_ch4_derived_co2e",
        "(derived at build)",
        "CH4-derived part of the upstream factor, per tonne LNG. The official "
        "inventory factor times (1 - upstream_ch4_share) is the CO2 part; any "
        "excess of the scenario upstream factor over that is methane-derived "
        "CO2e. This is the construction the upstream factor is built from, not "
        "a new assumption.",
        "tCO2e per t LNG",
        "derived",
        "max(upstream factor for the scenario - inventory_as_reported x "
        "(1 - upstream_ch4_share), 0). At measurement_central: "
        "0.25 - 0.22 x 0.7 = 0.096.",
    ),
    (
        "shipping_ch4_derived_co2e_share",
        "(derived at build)",
        "CH4-derived share of the shipping stage's CO2e. The 1.44 LNG carrier "
        "uplift in the shipping factor's own derivation is entirely measured "
        "methane slip, so that share of the shipping CO2e is methane rather "
        "than CO2.",
        "fraction",
        "derived",
        "1 - 1 / lng_carrier_methane_slip_uplift = 1 - 1/1.44 = 0.3056. "
        "At the 0.12 central shipping factor that is 0.0367 tCO2e per t LNG.",
    ),
    (
        "ch4_mass_kt",
        "(derived at build)",
        "Methane mass behind the CH4-derived CO2e in an asset-year. Not defined "
        "for the near_term_methane_gwp20 scenario, whose upstream factor "
        "(0.22 x 1.5) is not a GWP100 CO2e figure; that scenario's cell is left "
        "blank rather than divided by the wrong GWP.",
        "kt CH4",
        "derived",
        "ch4_derived_co2e_mt / gwp100_ch4 x 1000. Pipeline fugitive methane is "
        "not split and is not in this mass; liquefaction, regasification and "
        "combustion are treated as CO2.",
    ),
]


def main() -> None:
    wb = openpyxl.load_workbook(BOOK)
    ws = wb[SHEET]
    existing = {
        str(ws.cell(row=r, column=1).value).strip()
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=1).value is not None
    }
    row = ws.max_row + 1
    added = []
    for values in ROWS:
        if values[0] in existing:
            print(f"skip (already present): {values[0]}")
            continue
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        added.append(values[0])
        row += 1
    if added:
        wb.save(BOOK)
        print(f"added to '{SHEET}': {', '.join(added)}")
    else:
        print("nothing to add")


if __name__ == "__main__":
    main()
