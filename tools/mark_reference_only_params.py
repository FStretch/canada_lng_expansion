"""Mark every Parameters row live or reference-only, and wire the dead ones.

One-off. No value changes. The model never writes Inputs/; this script does, once.

Twenty-three rows on the Parameters sheet are read by no module. They are not
mistakes: route distances, coal comparators, the shipping reconstruction and the
target pathway are the context a reviewer needs to answer "where did this come
from". Deleting them would save nothing and cost traceability.

What this does instead, following the pattern already set by
`post_fid_to_completion`:

1. Adds a `status` column, `live` or `reference-only`. The value is DERIVED, not
   typed by hand: the script greps src/*.py and build_results.py for the quoted
   parameter name, so the column cannot drift away from the code.
2. Gives every reference-only row a one-line note saying what it is for and, where
   one exists, naming its live counterpart explicitly, so the workbook never
   carries two disconnected versions of one fact.
3. Records the new column on the workbook README sheet.

Three rows have no live counterpart at all - `coal_plant_reference_mw`,
`coal_plant_capacity_factor`, `coal_tco2e_per_mwh`. They support a coal-plant
equivalence comparator that is in no figure, no table and no output. Their notes
say so plainly rather than implying a connection that does not exist.

Run from the repository root:

    python tools/mark_reference_only_params.py
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "Inputs" / "Canada_LNG_Data_Inputs.xlsx"

# One line per reference-only row: what it is for, and its live counterpart.
NOTES = {
    "post_fid_to_completion": None,  # already done, left exactly as it stands
    "methane_leakage_low": (
        "REFERENCE-ONLY. Read by no module. An earlier three-point band for the "
        "multiplier by which measurement-based upstream methane exceeds the "
        "reported inventory. LIVE COUNTERPART: upstream_measurement_correction "
        "= 1.5 (cited), which is the single multiplier the model applies, and "
        "the Scenarios sheet, which carries the actual upstream range as whole "
        "factors (0.22 inventory_as_reported, 0.25 measurement_central, 0.26 "
        "measurement_high, 0.55 howarth_high). Do not cite this band and "
        "upstream_measurement_correction as two independent facts; they are two "
        "versions of one."
    ),
    "methane_leakage_mid": (
        "REFERENCE-ONLY. Read by no module. See methane_leakage_low. LIVE "
        "COUNTERPART: upstream_measurement_correction = 1.5. This 1.25 is the "
        "superseded central; the model uses 1.5."
    ),
    "methane_leakage_high": (
        "REFERENCE-ONLY. Read by no module. See methane_leakage_low. LIVE "
        "COUNTERPART: upstream_measurement_correction = 1.5, and the "
        "howarth_high scenario for the genuine upper bracket."
    ),
    "lng_to_gas_bcm_per_mtpa": None,  # already carries its own explanation
    "canada_2005_baseline": (
        "REFERENCE-ONLY. Read by no module. The reference year against which "
        "Canada's 2030 and 2035 targets are expressed, so it is what makes the "
        "target numbers interpretable. LIVE COUNTERPARTS: "
        "canada_2030_target_high and canada_2035_target_high, read by "
        "src/trajectories.py for the target pathway in figures 4, 5 and 6."
    ),
    "canada_2030_projected": (
        "REFERENCE-ONLY. Read by no module. The government's own projection of "
        "where 2030 emissions actually land, against a target of 417-455. Kept "
        "because it is the evidence for the statement that Canada is not on "
        "track. LIVE COUNTERPART: canada_2030_overshoot_gap = 200, which is the "
        "figure the model reads for the overshoot line."
    ),
    "canada_2035_target_low": (
        "REFERENCE-ONLY. Read by no module. The low end of the 2035 target "
        "range. LIVE COUNTERPART: canada_2035_target_high = 417, read by "
        "src/trajectories.py:644 for the target pathway. The high end is "
        "plotted because it is the less demanding test of the LNG buildout."
    ),
    "coal_plant_reference_mw": (
        "REFERENCE-ONLY, AND GENUINELY ORPHANED. Read by no module, and it has "
        "NO live counterpart: the coal-plant equivalence comparator these three "
        "rows would support is in no figure, no slide table and no output. Kept "
        "only so the three rows are not silently dropped. See also "
        "coal_plant_capacity_factor and coal_tco2e_per_mwh. If a reviewer asks "
        "what these are for, the honest answer is: nothing, at present."
    ),
    "coal_plant_capacity_factor": (
        "REFERENCE-ONLY, AND GENUINELY ORPHANED. See coal_plant_reference_mw. "
        "No live counterpart; supports no comparator now in the model."
    ),
    "coal_tco2e_per_mwh": (
        "REFERENCE-ONLY, AND GENUINELY ORPHANED. See coal_plant_reference_mw. "
        "No live counterpart; supports no comparator now in the model."
    ),
    "bc_lng_emission_intensity_benchmark": (
        "REFERENCE-ONLY. Read by no module. The British Columbia LNG emission "
        "intensity benchmark, kept as the yardstick a reviewer will hold the "
        "liquefaction factor against. LIVE COUNTERPARTS: the Emission Factors "
        "liquefaction central of 0.29 tCO2e/t LNG, which exceeds this benchmark "
        "by 0.13 (about 1.8 Mt/yr across 14 mtpa), and the electric-drive "
        "sensitivity in src/drive_sensitivity.py, which is the case in which a "
        "facility could approach it."
    ),
    "route_us_gulf_to_northeast_asia_nm": (
        "REFERENCE-ONLY. Read by no module. The comparator voyage a reviewer "
        "reaches for when asking whether a British Columbia route is short. "
        "LIVE COUNTERPART: route_bc_to_northeast_asia_nm = 3800, and the "
        "route_distance_nm column on the Asset Register, which is what "
        "src/model.route_scale_factor actually reads to scale shipping per "
        "asset."
    ),
    "voyage_days_bc_to_north_asia": (
        "REFERENCE-ONLY. Read by no module. Voyage duration for the British "
        "Columbia to north-east Asia run, kept as context for the shipping "
        "factor. LIVE COUNTERPART: route_bc_to_northeast_asia_nm = 3800 and the "
        "Emission Factors shipping central of 0.12; the model works in distance, "
        "not days."
    ),
    "ocean_transport_intensity_range_low": (
        "REFERENCE-ONLY as a Parameters row: read by no module HERE. It is not "
        "dead data. LIVE COUNTERPART: it is the same number as the Emission "
        "Factors shipping range_low of 0.05 tCO2e/t, which the Monte Carlo does "
        "sample. This row is the duplicate; the Emission Factors cell is the "
        "one that runs. Change that cell, not this one."
    ),
    "ocean_transport_intensity_range_high": (
        "REFERENCE-ONLY as a Parameters row: read by no module HERE. LIVE "
        "COUNTERPART: the same number as the Emission Factors shipping "
        "range_high of 0.31 tCO2e/t, which the Monte Carlo samples. This row is "
        "the duplicate; the Emission Factors cell is the one that runs."
    ),
    "route_atlantic_routing_factor": (
        "REFERENCE-ONLY. Read by no module. The great-circle-to-sailing-distance "
        "uplift used when the Atlantic route distances below were derived, "
        "calibrated on route_saint_john_to_rotterdam_nm. LIVE COUNTERPART: the "
        "route_distance_nm column on the Asset Register, which holds the "
        "finished distances this factor was used to produce; that column is what "
        "src/model.route_scale_factor reads. This row is the working, not the "
        "result."
    ),
    "route_baie_comeau_to_rotterdam_nm": (
        "REFERENCE-ONLY. Read by no module. Derived with "
        "route_atlantic_routing_factor; an estimate, not a published sailing "
        "distance. LIVE COUNTERPART: this asset's route_distance_nm cell on the "
        "Asset Register, which is the value the model reads."
    ),
    "route_fermeuse_to_rotterdam_nm": (
        "REFERENCE-ONLY. Read by no module. Derived with "
        "route_atlantic_routing_factor; an estimate, not a published sailing "
        "distance. LIVE COUNTERPART: this asset's route_distance_nm cell on the "
        "Asset Register."
    ),
    "route_saint_john_to_rotterdam_nm": (
        "REFERENCE-ONLY. Read by no module. The one cited Atlantic sailing "
        "distance, and the calibration point from which "
        "route_atlantic_routing_factor was derived. LIVE COUNTERPART: this "
        "asset's route_distance_nm cell on the Asset Register."
    ),
    "lng_carrier_aer_gco2_per_dwt_nm": (
        "REFERENCE-ONLY. Read by no module. The IMO Data Collection System fleet "
        "figure behind the independent reconstruction of the shipping factor: "
        "8.50 gCO2/DWT/nm over a 7,600 nm round trip, converted to cargo tonnes "
        "with cargo_tonnes_per_dwt and uplifted 44 per cent for measured methane "
        "slip, giving 0.110. LIVE COUNTERPART: the Emission Factors shipping "
        "central of 0.12 tCO2e/t, which this reconstruction independently "
        "supports and which its range_sources cell describes in full."
    ),
    "lng_carrier_kgco2_per_nm": (
        "REFERENCE-ONLY. Read by no module. Per ship, not per tonne of cargo, so "
        "it cannot be used as an intensity without a cargo denominator. Kept as "
        "the absolute-terms cross-check on "
        "lng_carrier_aer_gco2_per_dwt_nm. LIVE COUNTERPART: the Emission Factors "
        "shipping central of 0.12 tCO2e/t."
    ),
    "cargo_tonnes_per_dwt": (
        "REFERENCE-ONLY. Read by no module. The cargo-to-deadweight ratio that "
        "converts lng_carrier_aer_gco2_per_dwt_nm into an intensity per tonne of "
        "LNG in the README shipping reconstruction. Still uncited, and listed in "
        "Outputs/CITATIONS_WANTED.md as such. LIVE COUNTERPART: the Emission "
        "Factors shipping central of 0.12 tCO2e/t. Because no module reads this "
        "row, the open citation affects the reconstruction narrative only, not "
        "any published number."
    ),
}

README_HEADING = "PARAMETER STATUS"
README_ROW = (
    "live or reference-only",
    "Whether the model reads this row. `live` means at least one module in "
    "src/ or build_results.py looks the parameter up by name. `reference-only` "
    "means no module reads it: the row is documentation, context or the working "
    "behind a value that lives elsewhere. Every reference-only row carries a "
    "note saying what it is for and, where one exists, naming its live "
    "counterpart, so the workbook never holds two disconnected versions of one "
    "fact. Added 3 September 2026; the column is derived by "
    "tools/mark_reference_only_params.py by scanning the code, not typed by "
    "hand.",
)


def _code_text() -> str:
    parts = [(ROOT / "build_results.py").read_text(encoding="utf-8", errors="replace")]
    for path in sorted((ROOT / "src").glob("*.py")):
        parts.append(path.read_text(encoding="utf-8", errors="replace"))
    return "\n".join(parts)


def _is_read(name: str, code: str) -> bool:
    escaped = re.escape(name)
    return re.search(rf'"{escaped}"|\'{escaped}\'', code) is not None


def _headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(row=1, column=c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=c).value is not None
    }


def main() -> None:
    code = _code_text()
    wb = openpyxl.load_workbook(DATA)
    ws = wb["Parameters"]
    headers = _headers(ws)

    if "status" in headers:
        status_col = headers["status"]
        print("status column already present")
    else:
        status_col = ws.max_column + 1
        ws.cell(row=1, column=status_col, value="status")
        print(f"status column added at column {status_col}")

    notes_col = headers["notes"]
    name_col = headers["parameter"]

    live = 0
    ref = 0
    noted = 0
    unhandled = []
    for r in range(2, ws.max_row + 1):
        raw = ws.cell(row=r, column=name_col).value
        if raw is None or not str(raw).strip():
            continue
        name = str(raw).strip()
        if _is_read(name, code):
            ws.cell(row=r, column=status_col, value="live")
            live += 1
            continue
        ws.cell(row=r, column=status_col, value="reference-only")
        ref += 1
        if name not in NOTES:
            unhandled.append(name)
            continue
        note = NOTES[name]
        if note is None:
            continue
        ws.cell(row=r, column=notes_col, value=note)
        noted += 1

    if unhandled:
        raise SystemExit(
            "reference-only rows with no note written: " + ", ".join(unhandled)
        )

    rd = wb["README"]
    rd_max = rd.max_row
    already = any(
        str(rd.cell(row=r, column=1).value).strip() == README_HEADING
        for r in range(1, rd_max + 1)
    )
    if already:
        print("README section for parameter status already present")
    else:
        rd.cell(row=rd_max + 2, column=1, value=README_HEADING)
        rd.cell(row=rd_max + 3, column=1, value=README_ROW[0])
        rd.cell(row=rd_max + 3, column=2, value=README_ROW[1])
        print("README sheet: PARAMETER STATUS section added")

    wb.save(DATA)
    print(f"live {live}   reference-only {ref}   notes written {noted}")
    print("saved")


if __name__ == "__main__":
    main()
