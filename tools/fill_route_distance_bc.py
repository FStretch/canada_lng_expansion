"""Task 4: fill the two blank `route_distance_nm` cells on the Asset Register.

Shipping intensity is about to be scaled by route distance per asset, and the
model fails loudly where the distance is blank on an asset whose chain includes
the shipping stage. Two headline-scope export assets were blank:

    discovery_t1t4  Campbell River, British Columbia
    kanata_lng      Prince Rupert, British Columbia

Both are Pacific-coast British Columbia terminals, so both take the same cited
west-coast-Canada basis already used for the other six BC export assets
(CAPP / Oxford Institute for Energy Studies, approximately 3,800 nm to Tokyo).
That figure is a coast-level basis, not a port-specific sailing distance; the
Asset Sources cell says so and a Data Gaps row records it.

No emission factor or parameter changes. Run once, from the repository root:

    python tools/fill_route_distance_bc.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

BOOK = Path(__file__).resolve().parents[1] / "Inputs" / "Canada_LNG_Asset_Register.xlsx"

DISTANCE_NM = 3800

SOURCE_TEXT = (
    "CAPP, The Case for Canadian LNG (Apr 2026), citing Oxford Institute for "
    "Energy Studies: https://www.capp.ca/wp-content/uploads/2026/05/"
    "The-Case-for-Canadian-LNG-April-17-2026.pdf - west coast Canada to Tokyo "
    "approximately 3,800 nautical miles. Applied here as the west-coast-Canada "
    "basis, not a port-specific sailing distance: no published port-to-port "
    "distance was located for this terminal. Same basis as the other British "
    "Columbia export assets in the register."
)

TARGETS = ("discovery_t1t4", "kanata_lng")

GAP_ROW = (
    "route_distance_nm",
    "Discovery LNG, Kanata LNG",
    "Filled at the 3,800 nm west-coast basis",
    "Neither proponent publishes a port-to-port sailing distance. Both are "
    "Pacific-coast British Columbia terminals, so both take the cited "
    "west-coast Canada to Tokyo figure used for the other BC export assets. "
    "Prince Rupert is nearer north Asia than Kitimat and Campbell River is "
    "further, so the coast-level basis is slightly high for Kanata and "
    "slightly low for Discovery.",
    "A published port-to-port sailing distance for each terminal.",
    "low - shipping is about 3% of lifecycle emissions and the two errors "
    "point in opposite directions",
)


def _col_index(ws, name: str) -> int:
    for col in range(1, ws.max_column + 1):
        if str(ws.cell(row=1, column=col).value).strip() == name:
            return col
    raise KeyError(f"Column '{name}' not found on sheet '{ws.title}'.")


def _row_index(ws, project_id: str) -> int:
    pid_col = _col_index(ws, "project_id")
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=pid_col).value).strip() == project_id:
            return row
    raise KeyError(f"project_id '{project_id}' not found on sheet '{ws.title}'.")


def main() -> None:
    wb = openpyxl.load_workbook(BOOK)

    reg = wb["Asset Register"]
    src = wb["Asset Sources"]
    reg_col = _col_index(reg, "route_distance_nm")
    src_col = _col_index(src, "route_distance_nm")

    changed = []
    for pid in TARGETS:
        r = _row_index(reg, pid)
        current = reg.cell(row=r, column=reg_col).value
        if current not in (None, "", "Not available"):
            print(f"skip (already set): {pid} = {current}")
            continue
        reg.cell(row=r, column=reg_col, value=DISTANCE_NM)
        src.cell(row=_row_index(src, pid), column=src_col, value=SOURCE_TEXT)
        changed.append(pid)

    gaps = wb["Data Gaps"]
    have = {
        (
            str(gaps.cell(row=r, column=1).value).strip(),
            str(gaps.cell(row=r, column=2).value).strip(),
        )
        for r in range(2, gaps.max_row + 1)
    }
    if (GAP_ROW[0], GAP_ROW[1]) not in have:
        row = gaps.max_row + 1
        for col, value in enumerate(GAP_ROW, start=1):
            gaps.cell(row=row, column=col, value=value)
        print("added Data Gaps row for route_distance_nm")

    wb.save(BOOK)
    print(f"route_distance_nm = {DISTANCE_NM} set for: {', '.join(changed) or 'none'}")


if __name__ == "__main__":
    main()
