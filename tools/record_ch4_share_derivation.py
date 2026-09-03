"""Record the upstream_ch4_share derivation from Johnson et al. (2023).

One-off. Does NOT change the value: 0.30 stays, still typed `assumption`.
The research brackets the share rather than pinning it, and moving it is a
paper decision with consequences the authors should take deliberately.

Source, read from the primary document 3 September 2026:
Johnson, M.R., Conrad, B.M., Tyner, D.R. (2023), "Creating measurement-based
oil and gas sector methane inventories using source-resolved aerial surveys",
Communications Earth & Environment 4, doi 10.1038/s43247-023-00769-7.
Open PDF: https://www.nature.com/articles/s43247-023-00769-7.pdf

Two derivations, both starting from that paper's British Columbia figures.
They bracket the share at roughly 18-34% and centre near 25%. The assumed
0.30 sits in the upper part of that bracket without being excluded.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "Inputs" / "Canada_LNG_Data_Inputs.xlsx"

JOHNSON = "https://doi.org/10.1038/s43247-023-00769-7"

SOURCE = (
    "Share of upstream oil and gas emissions that is methane rather than "
    "carbon dioxide. Required to apply the measurement correction and the "
    "GWP20 uplift to the methane portion only. ASSUMPTION, DECLARED - 0.30 is "
    "not a cited figure. It is now bracketed by a derivation from Johnson, "
    "Conrad & Tyner (2023), Communications Earth & Environment, "
    + JOHNSON
    + " , which reports for British Columbia in 2021: total measurement-based "
    "upstream oil and gas methane of 144.5 kt/y; an upstream methane intensity "
    "of marketed natural gas of 0.38% (95% CI 0.33-0.44%, SI section S3); and "
    "that this measurement-based inventory is nominally 1.7 times the official "
    "ECCC bottom-up figure (range 1.5-2.0). "
    "ROUTE 1, intensity: dividing 0.38% by 1.7 gives an official-basis "
    "intensity of 0.224%; per tonne of marketed gas that is 0.00224 t CH4, "
    "which at GWP100 29.8 is 0.0666 tCO2e, or 30.3% of the 0.22 "
    "inventory_as_reported upstream factor. "
    "ROUTE 2, absolute: 144.5 / 1.7 = 85.0 kt CH4 official basis = 2.53 "
    "MtCO2e at 29.8; the 0.22 factor over 63 bcm (45.7 Mt LNG-equivalent at "
    "1.38 bcm per mtpa) implies 10.04 MtCO2e of upstream emissions net of "
    "transmission, so the share is 25.2%. "
    "BRACKET: across the 1.5-2.0 correction range and AR5 (25) versus AR6 "
    "(29.8) methane GWP vintages, the two routes span roughly 18% to 34% and "
    "centre near 25%. The derivation therefore CORROBORATES 0.30 as being "
    "within range while suggesting it sits high. It does not pin the value, so "
    "0.30 is retained and stays typed assumption."
)

NOTES = (
    "DO NOT substitute the national fugitives share. ECCC's 2026 executive "
    "summary reports 46 MtCO2e of oil and gas fugitive methane within roughly "
    "66 MtCO2e of total Fugitive Sources, a share near 70%. This model's "
    "upstream stage covers production and processing in total, INCLUDING "
    "stationary combustion and flaring, which are nearly all CO2. Applying a "
    "fugitives-only share here would overstate methane by roughly a factor of "
    "two. Canada's NIR gives about 20% for the oil and gas sector nationally, "
    "but that includes oil sands, which are far more CO2-intensive than gas "
    "production. "
    "SENSITIVITY OF THE MODEL TO THIS VALUE: the central upstream factor is "
    "insensitive - 0.22 x (1 - s + 1.5s) gives 0.2530 at s=0.30 and 0.2475 at "
    "s=0.25, both rounding to 0.25 - and measurement_high likewise rounds to "
    "0.26 either way. What moves is the near_term_methane_gwp20 scenario "
    "(0.428 at s=0.30 against 0.393 at s=0.25) and the CO2/CH4 split in the "
    "physics, which flows into the SC-CH4 damages line. "
    "OTHER ROUTES ATTEMPTED 3 September 2026 and not closed: the ECCC National "
    "Inventory Report Part 3 provincial tables (category 1.B.2 for British "
    "Columbia) exist but are served through a JavaScript data-mart portal "
    "rather than a direct download, so the CO2/CH4 cells could not be pulled; "
    "the UNFCCC Common Reporting Tables for Canada are the route that would "
    "work. The BC Provincial Inventory methodology report confirms upstream "
    "fugitives are Tier 3 bottom-up facility by facility, so the underlying "
    "data distinguishes gases even though the published provincial tables are "
    "CO2e only."
)


def main() -> None:
    wb = openpyxl.load_workbook(DATA)
    ws = wb["Parameters"]
    headers = {
        str(ws.cell(row=1, column=c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=c).value is not None
    }
    row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=headers["parameter"]).value).strip() == "upstream_ch4_share":
            row = r
            break
    if row is None:
        raise KeyError("upstream_ch4_share not found")

    value = ws.cell(row=row, column=headers["value"]).value
    ws.cell(row=row, column=headers["source"], value=SOURCE)
    ws.cell(row=row, column=headers["source_url"], value=JOHNSON)
    ws.cell(row=row, column=headers["notes"], value=NOTES)
    wb.save(DATA)
    print(f"upstream_ch4_share stays {value}; derivation and bracket recorded")


if __name__ == "__main__":
    main()
